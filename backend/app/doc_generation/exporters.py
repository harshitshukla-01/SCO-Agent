import re
from html import escape
from io import BytesIO
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from app.agent.policy_generator import DISCLAIMER
from app.models.policy import Policy, PolicyVersion


def _metadata_rows(
    policy: Policy,
    version: PolicyVersion,
    organization_name: str,
    approved_by: Optional[str],
    approved_at: Optional[str],
):
    version_status = version.status.value if version.status else policy.status.value
    return [
        ("Organization", organization_name),
        ("Policy", policy.name),
        ("Version", f"v{version.version_number}"),
        ("Status", version_status.replace("_", " ").title()),
        ("Owner", policy.owner or "Unassigned"),
        ("Approved by", approved_by or version.approved_by_email or "Not recorded"),
        ("Approval date", approved_at or version.approved_at or "Not recorded"),
        ("Next review date", policy.next_review or "Not scheduled"),
    ]


def _markdown_to_docx(document: Document, content: str):
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 3))
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", stripped):
            document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
        else:
            paragraph = document.add_paragraph()
            pieces = re.split(r"(\*\*.+?\*\*|`[^`]+`)", stripped)
            for piece in pieces:
                if piece.startswith("**") and piece.endswith("**"):
                    paragraph.add_run(piece[2:-2]).bold = True
                elif piece.startswith("`") and piece.endswith("`"):
                    run = paragraph.add_run(piece[1:-1])
                    run.font.name = "Consolas"
                else:
                    paragraph.add_run(piece)


def build_docx(
    policy: Policy,
    version: PolicyVersion,
    organization_name: str,
    approved_by: Optional[str],
    approved_at: Optional[str],
    is_draft: bool,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if is_draft:
        header = section.header.paragraphs[0]
        header.text = "DRAFT - NOT APPROVED"
        header.alignment = 1
        header.runs[0].bold = True
        header.runs[0].font.color.rgb = RGBColor(170, 25, 25)
    title = document.add_heading(policy.name, level=0)
    title.paragraph_format.space_after = Pt(6)
    metadata = document.add_table(rows=0, cols=2)
    metadata.style = "Light Shading Accent 1"
    for label, value in _metadata_rows(policy, version, organization_name, approved_by, approved_at):
        cells = metadata.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True

    document.add_paragraph()
    _markdown_to_docx(document, version.content)
    footer = section.footer.paragraphs[0]
    footer.text = DISCLAIMER
    footer.alignment = 1
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_markdown_paragraph(text: str) -> str:
    escaped = escape(text)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", escaped)
    return escaped


def build_pdf(
    policy: Policy,
    version: PolicyVersion,
    organization_name: str,
    approved_by: Optional[str],
    approved_at: Optional[str],
    is_draft: bool,
) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PolicyTitle", parent=styles["Title"], alignment=TA_LEFT, textColor=colors.HexColor("#17324d"), spaceAfter=10))
    styles.add(ParagraphStyle(name="PolicyBody", parent=styles["BodyText"], leading=15, spaceAfter=7))
    styles.add(ParagraphStyle(name="PolicyBullet", parent=styles["BodyText"], leftIndent=18, firstLineIndent=-10, leading=15, spaceAfter=4))
    styles.add(ParagraphStyle(name="PolicyMeta", parent=styles["BodyText"], fontSize=9, leading=13))
    story = [Paragraph(escape(policy.name), styles["PolicyTitle"])]
    for label, value in _metadata_rows(policy, version, organization_name, approved_by, approved_at):
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", styles["PolicyMeta"]))
    story.extend([
        Spacer(1, 10),
        HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#9aafc2")),
        Spacer(1, 14),
    ])

    for line in version.content.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 5))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            style = styles["Heading1"] if level == 1 else styles["Heading2"] if level == 2 else styles["Heading3"]
            story.append(Paragraph(_pdf_markdown_paragraph(heading.group(2)), style))
        elif stripped.startswith(("- ", "* ")):
            story.append(Paragraph("&bull; " + _pdf_markdown_paragraph(stripped[2:]), styles["PolicyBullet"]))
        elif re.match(r"^\d+[.)]\s+", stripped):
            story.append(Paragraph(_pdf_markdown_paragraph(stripped), styles["PolicyBullet"]))
        else:
            story.append(Paragraph(_pdf_markdown_paragraph(stripped), styles["PolicyBody"]))

    def decorate(canvas, doc):
        canvas.saveState()
        if is_draft:
            canvas.setFillColor(colors.Color(0.75, 0.1, 0.1, alpha=0.14))
            canvas.translate(letter[0] / 2, letter[1] / 2)
            canvas.rotate(35)
            canvas.setFont("Helvetica-Bold", 34)
            canvas.drawCentredString(0, 0, "DRAFT - NOT APPROVED")
            canvas.rotate(-35)
            canvas.translate(-letter[0] / 2, -letter[1] / 2)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawCentredString(letter[0] / 2, 0.35 * inch, DISCLAIMER)
        canvas.drawRightString(letter[0] - 0.65 * inch, 0.35 * inch, str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.7 * inch,
        title=policy.name,
        author=organization_name,
    ).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()