from docx import Document
from io import BytesIO
import pytest
from fastapi import HTTPException

import app.api.admin.policy_generation as policy_generation_api_module
import app.agent.policy_generator as policy_generator_module
from app.models.project_facts import FactSource, ProjectFactValue, ProjectFacts


ADMIN_A = {"Authorization": "Bearer admin-token-org-a"}
USER_A = {"Authorization": "Bearer user-token-org-a"}
ADMIN_B = {"Authorization": "Bearer admin-token-org-b"}
USER_B = {"Authorization": "Bearer user-token-org-b"}


def _make_published_policy(client):
    response = client.post(
        "/api/admin/policies",
        json={
            "name": "Access Control Policy",
            "type": "technical",
            "initial_content": "# Existing published policy\n\nCurrent approved content.",
        },
        headers=ADMIN_A,
    )
    assert response.status_code == 201
    policy = response.json()
    for target in ("in_review", "approved", "published"):
        response = client.patch(
            f"/api/admin/policies/{policy['id']}/status",
            json={"status": target},
            headers=ADMIN_A,
        )
        assert response.status_code == 200
    return policy["id"]


def test_generation_falls_back_to_local_template_when_gemini_key_missing(monkeypatch):
    facts = ProjectFacts(
        facts={
            "company_name": ProjectFactValue(
                value="Example Co",
                source=FactSource.MANUAL,
                evidence="Confirmed by admin",
                updated_at="2026-01-01T00:00:00Z",
            )
        }
    )

    class FakeSettings:
        GEMINI_API_KEY = None
        GEMINI_MODEL = "gemini-2.5-flash"

    monkeypatch.setattr(policy_generator_module, "get_settings", lambda: FakeSettings())

    content, model_name, fact_sources = policy_generator_module.generate_policy_content(
        "Access Control Policy",
        "pol-access-control-policy",
        "technical",
        facts,
    )

    assert model_name == "local-template"
    assert "Access Control Policy" in content
    assert "AI-generated draft" in content
    assert fact_sources and fact_sources[0]["field"] == "company_name"


def test_generation_keeps_published_version_until_candidate_is_published(client, monkeypatch):
    policy_id = _make_published_policy(client)
    facts_response = client.put(
        "/api/admin/project-facts",
        json={"fields": {"company_name": {"value": "Example Co", "evidence": "Confirmed by admin"}}},
        headers=ADMIN_A,
    )
    assert facts_response.status_code == 200
    assert facts_response.json()["version"] == 1

    def fake_generate(name, identifier, policy_type, facts):
        assert identifier == "pol-access-control-policy" or identifier == policy_id
        return (
            "# Access Control Policy\n\n## Purpose\nExample Co defines access requirements. [Fact: company_name]\n",
            "gemini-test-model",
            [{"field": "company_name", "value": "Example Co", "source": "manual", "evidence": "Confirmed by admin", "updated_at": "now"}],
        )

    monkeypatch.setattr(policy_generation_api_module, "generate_policy_content", fake_generate)
    generated = client.post(f"/api/admin/policies/{policy_id}/generate", headers=ADMIN_A)
    assert generated.status_code == 201
    generated_version = generated.json()
    assert generated_version["source"] == "ai"
    assert generated_version["status"] == "draft"
    assert generated_version["fact_sheet_version"] == 1
    assert generated_version["model_name"] == "gemini-test-model"
    generation_audit = client.get("/api/admin/audit-log", headers=ADMIN_A).json()
    assert any(
        entry["action"] == "POLICY_AI_DRAFT_GENERATED"
        and entry["details"]["fact_sheet_version"] == 1
        for entry in generation_audit
    )

    edited_content = "# Access Control Policy\n\n## Purpose\nEdited policy text. [Fact: company_name]\n"
    edited = client.post(
        f"/api/admin/policies/{policy_id}/versions",
        json={"content": edited_content, "change_note": "Admin edits", "derived_from_version": generated_version["version_number"]},
        headers=ADMIN_A,
    )
    assert edited.status_code == 201
    edited_version = edited.json()
    assert edited_version["source"] == "ai"
    assert edited_version["model_name"] == "gemini-test-model"
    assert edited_version["fact_sheet_version"] == 1
    assert edited_version["fact_sources"][0]["field"] == "company_name"

    policy = client.get(f"/api/admin/policies/{policy_id}", headers=ADMIN_A).json()
    assert policy["status"] == "published"
    assert policy["current_version"] == 1
    assert policy["draft_version"] == 3
    assert policy["latest_version"]["content"].startswith("# Existing published policy")

    member_policy = client.get(f"/api/user/policies/{policy_id}", headers=USER_A)
    assert member_policy.status_code == 200
    assert member_policy.json()["latest_version"]["content"].startswith("# Existing published policy")

    for target in ("in_review", "approved", "published"):
        transitioned = client.patch(
            f"/api/admin/policies/{policy_id}/versions/3/status",
            json={"status": target},
            headers=ADMIN_A,
        )
        assert transitioned.status_code == 200
        if target != "published":
            unchanged = client.get(f"/api/admin/policies/{policy_id}", headers=ADMIN_A).json()
            assert unchanged["current_version"] == 1

    published = client.get(f"/api/user/policies/{policy_id}", headers=USER_A).json()
    assert published["current_version"] == 3
    assert published["latest_version"]["content"] == edited_content
    pdf = client.get(f"/api/user/policies/{policy_id}/export/pdf?version=3", headers=USER_A)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")
    assert pdf.content.rstrip().endswith(b"%%EOF")
    export_audit = client.get("/api/admin/audit-log", headers=ADMIN_A).json()
    assert any(
        entry["action"] == "POLICY_EXPORTED"
        and entry["details"]["version"] == 3
        and entry["details"]["format"] == "pdf"
        for entry in export_audit
    )


def test_user_cannot_generate_and_facts_versions_are_tenant_scoped(client, monkeypatch):
    policy_id = client.post(
        "/api/admin/policies",
        json={"name": "Org A Policy", "type": "technical"},
        headers=ADMIN_A,
    ).json()["id"]

    denied = client.post(f"/api/admin/policies/{policy_id}/generate", headers=USER_A)
    assert denied.status_code == 403

    saved = client.put(
        "/api/admin/project-facts",
        json={"fields": {"company_name": {"value": "Org A", "evidence": "Test"}}},
        headers=ADMIN_A,
    )
    assert saved.status_code == 200
    other_facts = client.get("/api/admin/project-facts", headers=ADMIN_B).json()
    assert other_facts["version"] == 0
    assert other_facts["facts"] == {}

    monkeypatch.setattr(
        policy_generation_api_module,
        "generate_policy_content",
        lambda *args: ("# Draft\n", "gemini-test-model", []),
    )
    generated = client.post(f"/api/admin/policies/{policy_id}/generate", headers=ADMIN_A)
    assert generated.status_code == 201
    assert client.get(f"/api/admin/policies/{policy_id}", headers=ADMIN_B).status_code == 404
    assert client.get(f"/api/admin/policies/{policy_id}/export/pdf", headers=ADMIN_B).status_code == 404
    assert client.get(f"/api/user/policies/{policy_id}/export/pdf", headers=USER_B).status_code == 404


def test_draft_export_requires_explicit_override_and_contains_marking(client):
    policy_id = client.post(
        "/api/admin/policies",
        json={"name": "Draft Export Policy", "type": "documentary", "initial_content": "# Draft\n\nReview this."},
        headers=ADMIN_A,
    ).json()["id"]

    denied = client.get(f"/api/admin/policies/{policy_id}/export/pdf", headers=ADMIN_A)
    assert denied.status_code == 403

    pdf = client.get(f"/api/admin/policies/{policy_id}/export/pdf?allow_draft=true", headers=ADMIN_A)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")

    docx = client.get(f"/api/admin/policies/{policy_id}/export/docx?allow_draft=true", headers=ADMIN_A)
    assert docx.status_code == 200
    document = Document(BytesIO(docx.content))
    assert "Draft Export Policy" in document.paragraphs[0].text
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    document_text += "\n" + "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    for field in ("Organization", "Version", "Status", "Owner", "Approved by", "Approval date", "Next review date"):
        assert field in document_text
    assert "AI-generated draft" in document.sections[0].footer.paragraphs[0].text
    assert "DRAFT - NOT APPROVED" in document.sections[0].header.paragraphs[0].text


def test_generation_rate_limit_is_enforced():
    key = ("rate-limit-org", "rate-limit-admin")
    policy_generation_api_module._generation_requests.pop(key, None)
    try:
        for _ in range(policy_generation_api_module._GENERATION_LIMIT):
            policy_generation_api_module._guard_generation_rate(*key)
        with pytest.raises(HTTPException) as error:
            policy_generation_api_module._guard_generation_rate(*key)
        assert error.value.status_code == 429
    finally:
        policy_generation_api_module._generation_requests.pop(key, None)